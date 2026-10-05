import json
import math
import re
from pathlib import Path
from openpyxl import load_workbook


def clean_code(value, number_format="General") -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        format_section = number_format.split(";", 1)[0]
        decimal_format = re.search(r"\.([0#?]+)", format_section)
        if decimal_format:
            return format(value, f".{len(decimal_format.group(1))}f").upper()
        code = str(value)
        if code.endswith(".0"):
            code = code[:-2]
        if re.fullmatch(r"0+", format_section):
            code = code.zfill(len(format_section))
        return code.upper()
    return clean_text(value).replace(" ", "").upper()


def clean_text(value) -> str:
    if value is None:
        return ""
    text = re.sub(r"\bqunital(s?)\b", r"quintal\1", str(value).strip(), flags=re.IGNORECASE)
    unit_replacements = {
        "meter": "metre",
        "meters": "metres",
        "kg": "kg",
        "kgs": "kgs",
        "kilogram": "kilogram",
        "kilograms": "kilograms",
        "mt": "mt",
        "mts": "mts",
        "ton": "ton",
        "tons": "tons",
        "no": "no",
        "nos": "nos",
    }
    return ONE_QUANTITY_RE.sub(
        lambda match: f"1 {unit_replacements.get(match.group(1).casefold(), match.group(1))}",
        text,
    )


def parse_number(value):
    try:
        number = float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


PERCENT_OF_P_RE = re.compile(
    r"(\d+(?:\.\d+)?)\s*(?:%|percentages?)\s*of\s*P\b", re.IGNORECASE
)
ONE_QUANTITY_RE = re.compile(
    r"\b(?:one|1)\s+(quintals?|metres?|meters?|cum|sqm|kgs?|kilograms?|"
    r"tonnes?|tons?|mts?|nos?\.?|each|trees?|holes?)\b",
    re.IGNORECASE,
)

def is_main_item_code(code):
    if "SUB" in code:
        return True

    parts = code.split(".")
    return (
        len(parts) > 1
        and parts[0].isdigit()
        and all(part and part.replace("(", "").replace(")", "").isalnum() for part in parts[1:])
    )


def parse_rate_analysis(filepath, output_file=None):
    print(f"Reading Rate Analysis from {filepath}...")
    try:
        workbook = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        print(f"Error opening {filepath}: {e}")
        return

    ra_data = {}
    for sheet in workbook.worksheets:
        sheet_name = sheet.title
        sheet_parts = sheet_name.strip().upper().split()
        if len(sheet_parts) != 2 or sheet_parts[0] != "SH" or not sheet_parts[1].isdigit():
            continue

        current_item: str | None = None
        current_description = ""
        current_analysis = []
        current_p_value: float | None = None

        def save_current_item():
            if current_item is None or current_item in {"1.1.2", "1.1.8"}:
                return
            existing = ra_data.get(current_item)
            if existing is None or len(current_analysis) > len(existing["analysis"]):
                ra_data[current_item] = {
                    "code": current_item,
                    "description": current_description,
                    "analysis": list(current_analysis),
                }

        for row in sheet.iter_rows(min_col=1, max_col=7):
            cols = [clean_text(cell.value) for cell in row]
            while len(cols) < 7:
                cols.append("")
            code_val = clean_code(row[0].value, row[0].number_format)
            cols[0] = code_val
            desc_val = cols[1]
            c_code = code_val

            if c_code.lower() == "code" or desc_val.lower() == "description": continue
            if not code_val and not desc_val and not cols[5]: continue

            has_detail_values = any(cols[index] for index in range(2, 6))
            is_item_reference = "RATE AS PER ITEM" in desc_val.upper()
            if c_code and is_main_item_code(c_code) and not has_detail_values and not is_item_reference:
                save_current_item()
                if c_code in {"1.1.2", "1.1.8"}:
                    current_item = None
                    current_description = ""
                    current_analysis = []
                    current_p_value = None
                    continue
                current_item = c_code
                current_description = desc_val
                current_analysis = []
                current_p_value = None
                continue

            if current_item:
                if current_item.startswith("3.") and any(x in desc_val.lower() for x in ["water charges", "cpoh", "gst", "cess"]): continue
                if current_item == "5.9.3" and re.fullmatch(r"detail of cost\s+for\s+10 room", desc_val, re.IGNORECASE):
                    desc_val = "Detail of cost for 135 sqm"
                marker = cols[6].strip().upper()
                row_amount = parse_number(cols[5])
                percentage_of_p = PERCENT_OF_P_RE.search(desc_val)
                if marker == "P" and row_amount is not None:
                    current_p_value = row_amount
                elif percentage_of_p and current_p_value is not None:
                    percentage = float(percentage_of_p.group(1))
                    row_amount = current_p_value * percentage / 100
                    cols[5] = f"{row_amount:.2f}"
                current_analysis.append({ "code": code_val, "description": desc_val, "unit": cols[2], "quantity": cols[3], "rate": cols[4], "amount": cols[5], "mark": cols[6] })

        save_current_item()

    output_path = Path(output_file) if output_file else Path(__file__).resolve().parent / "public" / "data" / "rate_analysis_master.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open('w', encoding='utf-8') as f:
        json.dump(ra_data, f, indent=4, ensure_ascii=False, allow_nan=False)
    print(f"Success: rate_analysis_master.json generated with {len(ra_data)} items.")

if __name__ == "__main__":
    parse_rate_analysis("dsr_2023_master.xlsx")