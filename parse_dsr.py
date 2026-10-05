import json
import math
import re
from pathlib import Path
from openpyxl import load_workbook


SUBHEAD_ONE_LABEL = "SH 1: CARRIAGE OF MATERIALS"
HARDCODED_DSR_ITEMS = {
    "1.1.2": {
        "code": "1.1.2",
        "description": "Earth",
        "unit": "cum",
        "rate": 326.68,
        "isMain": False,
    },
    "1.1.8": {
        "code": "1.1.8",
        "description": "Bricks",
        "unit": "1000 Nos",
        "rate": 326.68,
        "isMain": False,
    },
}
ONE_QUANTITY_RE = re.compile(
    r"\b(?:one|1)\s+(quintals?|metres?|meters?|cum|sqm|kgs?|kilograms?|"
    r"tonnes?|tons?|mt|nos?\.?|each|trees?|holes?)\b",
    re.IGNORECASE,
)
ONE_QUANTITY_UNIT_REPLACEMENTS = {
    "meter": "metre",
    "meters": "metres",
}


def clean_text(value) -> str:
    if value is None:
        return ""
    text = re.sub(r"\bqunital(s?)\b", r"quintal\1", str(value).strip(), flags=re.IGNORECASE)
    return ONE_QUANTITY_RE.sub(
        lambda match: (
            "1 "
            + ONE_QUANTITY_UNIT_REPLACEMENTS.get(
                match.group(1).casefold(), match.group(1).casefold()
            )
        ),
        text,
    )


def clean_code(value, number_format="General") -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        format_section = number_format.split(";", 1)[0]
        decimal_format = re.search(r"\.([0#?]+)", format_section)
        if decimal_format:
            return format(value, f".{len(decimal_format.group(1))}f").upper()
        code = str(value)
        if code.endswith(".0"):
            code = code[:-2]
        if re.fullmatch(r"0+", format_section):
            code = code.zfill(len(format_section))
        return code.upper()
    return clean_text(value).replace(" ", "").upper()


def is_dsr_code(code):
    parts = code.split(".")
    return (
        len(parts) > 1
        and parts[0].isdigit()
        and all(part and part.replace("(", "").replace(")", "").isalnum() for part in parts[1:])
    )


def parse_dsr_stable(file_path, output_file=None):
    print(f"Reading {file_path}...")
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        print(f"Error opening file: {e}")
        return

    dsr_data = {}

    for sheet in workbook.worksheets:
        sheet_name = sheet.title
        print(f"Scanning sheet: {sheet_name}")
        if sheet.max_column < 4:
            continue

        for row in sheet.iter_rows(min_row=2, max_col=4):
            code = clean_code(row[0].value, row[0].number_format)
            if not is_dsr_code(code):
                continue

            description = clean_text(row[1].value)
            unit = clean_text(row[2].value)
            try:
                rate = float(str(row[3].value).replace(",", "").strip())
            except (TypeError, ValueError):
                rate = 0.0
            if not math.isfinite(rate):
                rate = 0.0

            chapter = code.split(".", 1)[0]
            if chapter == "1" and code not in HARDCODED_DSR_ITEMS:
                continue
            group_key = f"sh_{chapter}"
            label = SUBHEAD_ONE_LABEL if chapter == "1" else f"SH {chapter}"
            group = dsr_data.setdefault(group_key, {"label": label, "items": []})
            if code in HARDCODED_DSR_ITEMS:
                continue
            group["items"].append({
                "code": code,
                "description": description,
                "unit": unit,
                "rate": rate,
                "isMain": rate <= 0,
            })

    dsr_data["sh_1"] = {
        "label": SUBHEAD_ONE_LABEL,
        "items": [dict(item) for item in HARDCODED_DSR_ITEMS.values()],
    }

    output_path = Path(output_file) if output_file else Path(__file__).resolve().parent / "public" / "data" / "pure_dsr.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open('w', encoding='utf-8') as f:
        json.dump(dsr_data, f, indent=4, ensure_ascii=False, allow_nan=False)

    item_count = sum(len(group["items"]) for group in dsr_data.values())
    print(f"Success! Extracted {item_count} items and saved to {output_path}.")

if __name__ == "__main__":
    parse_dsr_stable("dsr_db_2.xlsx")